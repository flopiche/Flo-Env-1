#!/usr/bin/env python3
"""
Analyse de cohérence produits — images vs fiche
Usage : python analyse_coherence.py export_produits.xlsx
"""

import sys
import os
import time
import base64
import json
import re
import urllib.request
import urllib.error
from pathlib import Path

# ── Dépendances ───────────────────────────────────────────────────────────────
try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("❌ Installez openpyxl : pip install openpyxl")

try:
    import anthropic
except ImportError:
    sys.exit("❌ Installez le SDK Anthropic : pip install anthropic")

# ── Config ────────────────────────────────────────────────────────────────────
MODEL          = "claude-haiku-4-5-20251001"   # rapide + vision + économique
MAX_IMAGES     = 6                              # images max par produit
DELAY_BETWEEN  = 1.0                           # secondes entre appels API
IMAGE_TIMEOUT  = 15                            # timeout téléchargement image (s)
MAX_TOKENS     = 900

PROMPT_SYSTEM = """Tu es un expert en contrôle qualité de fiches produit e-commerce.
Tu analyses la cohérence entre une fiche produit et ses visuels.
Tu réponds UNIQUEMENT en JSON valide, sans markdown, sans commentaire."""

PROMPT_USER = """Analyse la cohérence entre cette fiche produit et ses images.

Fiche produit :
- Libellé : {title}
- Description : {desc}
- Couleur déclarée : {color}
- Tranche d'âge : {age}
- Âge minimum : {age_min}

Règle absolue : une info ABSENTE du visuel n'est PAS une anomalie.
Ne marque "contradiction" que si l'image CONTREDIT activement la fiche.

Verdicts possibles : "cohérent" | "contradiction" | "non vérifiable"

Évalue ces critères :
1. type_produit : le produit visible correspond-il au libellé/description ?
2. motif_design : le motif/design décrit est-il cohérent ?
3. couleur : la couleur visible correspond-elle à la couleur déclarée ?
4. contenu_set : le contenu visible correspond-il à ce qui est décrit ?
5. public_age : le visuel est-il cohérent avec la tranche d'âge ?
6. qualite_image : images nettes, bien cadrées, utilisables pour une fiche ?

Réponds en JSON strict :
{{
  "verdict_global": "cohérent|contradiction|non vérifiable",
  "type_produit":   {{"verdict": "...", "remarque": "..."}},
  "motif_design":   {{"verdict": "...", "remarque": "..."}},
  "couleur":        {{"verdict": "...", "remarque": "..."}},
  "contenu_set":    {{"verdict": "...", "remarque": "..."}},
  "public_age":     {{"verdict": "...", "remarque": "..."}},
  "qualite_image":  {{"verdict": "...", "remarque": "..."}},
  "remarques_generales": "..."
}}"""

# ── Couleurs Excel ────────────────────────────────────────────────────────────
FILL_GREEN  = PatternFill("solid", fgColor="D4EDDA")
FILL_RED    = PatternFill("solid", fgColor="F8D7DA")
FILL_GREY   = PatternFill("solid", fgColor="E2E3E5")
FILL_ORANGE = PatternFill("solid", fgColor="FFF3CD")
FILL_HEADER = PatternFill("solid", fgColor="343A40")
FONT_HEADER = Font(color="FFFFFF", bold=True, size=10)
FONT_BOLD   = Font(bold=True)

def verdict_fill(v):
    if v == "cohérent":      return FILL_GREEN
    if v == "contradiction":  return FILL_RED
    return FILL_GREY

# ── Lecture Excel ─────────────────────────────────────────────────────────────
def find_col(headers, terms):
    for t in terms:
        for i, h in enumerate(headers):
            if h and t.lower() in str(h).lower():
                return i
    return -1

def load_products(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((s for s in wb.sheetnames if s.lower() == "data"), wb.sheetnames[0])
    ws = wb[sheet]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if len(rows) < 2:
        sys.exit(f"❌ Feuille '{sheet}' vide ou illisible.")

    # Détecte double ligne d'en-tête
    headers = [str(h or "").strip() for h in rows[0]]
    data_start = 1
    if len(rows) >= 3:
        ean_col = find_col(headers, ["ean"])
        if ean_col != -1:
            val2 = str(rows[1][ean_col] or "").strip()
            if not re.match(r"^\d{8,}$", val2):
                data_start = 2  # row 2 est aussi un en-tête

    c_ean   = find_col(headers, ["ean", "code barre", "barcode"])
    c_title = find_col(headers, ["libellé", "libelle", "titre", "title", "fr-product-name", "nom"])
    c_desc  = find_col(headers, ["description", "fr-description"])
    c_color = find_col(headers, ["couleur principale", "couleur", "color", "main-color"])
    c_age   = find_col(headers, ["tranche", "âge", "age-range"])
    c_agemin= find_col(headers, ["age minimum", "âge minimum", "minimum-age"])

    img_cols = []
    c0 = find_col(headers, ["image principale", "main-image", "image 1", "image1"])
    if c0 != -1: img_cols.append(c0)
    for n in range(2, 9):
        c = find_col(headers, [f"image {n}", f"image{n}", f"image-{n}"])
        if c != -1: img_cols.append(c)

    products = []
    for row in rows[data_start:]:
        ean = str(row[c_ean] if c_ean != -1 else "").strip()
        if not re.match(r"^\d{8,}$", ean):
            continue
        products.append({
            "ean":     ean,
            "title":   str(row[c_title]  if c_title  != -1 else "").strip(),
            "desc":    str(row[c_desc]   if c_desc   != -1 else "").strip(),
            "color":   str(row[c_color]  if c_color  != -1 else "").strip(),
            "age":     str(row[c_age]    if c_age    != -1 else "").strip(),
            "age_min": str(row[c_agemin] if c_agemin != -1 else "").strip(),
            "img_urls": [
                str(row[c] or "").strip()
                for c in img_cols
                if c < len(row) and row[c] and str(row[c]).strip().startswith("http")
            ]
        })
    return products

# ── Téléchargement image ──────────────────────────────────────────────────────
def fetch_image_b64(url):
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=IMAGE_TIMEOUT) as r:
            data = r.read()
            ct   = r.headers.get_content_type() or "image/jpeg"
            return base64.standard_b64encode(data).decode(), ct
    except Exception:
        return None, None

# ── Appel Claude ──────────────────────────────────────────────────────────────
def analyze(client, product, images_b64):
    content = []
    for b64, ct in images_b64[:MAX_IMAGES]:
        content.append({
            "type": "image",
            "source": {"type": "base64", "media_type": ct, "data": b64}
        })
    if not images_b64:
        content.append({"type": "text", "text": "(Aucune image accessible pour ce produit)"})

    content.append({"type": "text", "text": PROMPT_USER.format(
        title   = product["title"]   or "(non renseigné)",
        desc    = (product["desc"]   or "(non renseignée)")[:800],
        color   = product["color"]   or "(non renseignée)",
        age     = product["age"]     or "(non renseignée)",
        age_min = product["age_min"] or "(non renseigné)",
    )})

    resp = client.messages.create(
        model=MODEL,
        max_tokens=MAX_TOKENS,
        system=PROMPT_SYSTEM,
        messages=[{"role": "user", "content": content}]
    )
    raw = resp.content[0].text.strip()
    raw = re.sub(r"^```json\s*", "", raw); raw = re.sub(r"```$", "", raw).strip()
    return json.loads(raw)

# ── Écriture Excel résultat ───────────────────────────────────────────────────
CRITERES = ["type_produit", "motif_design", "couleur", "contenu_set", "public_age", "qualite_image"]
CRITERES_FR = ["Type produit", "Motif/design", "Couleur", "Contenu set", "Public/âge", "Qualité image"]

def write_results(results, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cohérence"

    # En-têtes
    headers = ["EAN", "Libellé", "Verdict global"] + \
              [f"{l}\nverdict" for l in CRITERES_FR] + \
              [f"{l}\nremarque" for l in CRITERES_FR] + \
              ["Remarques générales", "Images analysées", "Erreur"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill   = FILL_HEADER
        cell.font   = FONT_HEADER
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    thin = Side(style="thin", color="DEE2E6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_i, r in enumerate(results, 2):
        res    = r.get("result", {})
        error  = r.get("error", "")
        vg     = res.get("verdict_global", "erreur" if error else "non vérifiable")

        row_data = [
            r["ean"],
            r["title"],
            vg,
        ] + [
            res.get(k, {}).get("verdict", "") for k in CRITERES
        ] + [
            res.get(k, {}).get("remarque", "") for k in CRITERES
        ] + [
            res.get("remarques_generales", ""),
            r.get("n_images", 0),
            error,
        ]

        for col_i, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_i, column=col_i, value=val)
            cell.border    = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")

            # Colorier verdict global
            if col_i == 3:
                cell.fill = verdict_fill(vg)
                cell.font = FONT_BOLD
            # Colorier verdicts critères
            elif 4 <= col_i <= 3 + len(CRITERES):
                cell.fill = verdict_fill(val)

    # Largeurs colonnes
    widths = [16, 30, 14] + [14]*len(CRITERES) + [28]*len(CRITERES) + [35, 8, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Figer les 3 premières colonnes
    ws.freeze_panes = "D2"

    # Onglet résumé
    ws2 = wb.create_sheet("Résumé")
    ws2.append(["Verdict", "Nombre", "%"])
    counts = {"cohérent": 0, "contradiction": 0, "non vérifiable": 0, "erreur": 0}
    for r in results:
        vg = r.get("result", {}).get("verdict_global", "erreur" if r.get("error") else "non vérifiable")
        counts[vg if vg in counts else "non vérifiable"] += 1
    total = len(results)
    fills = {"cohérent": FILL_GREEN, "contradiction": FILL_RED, "non vérifiable": FILL_GREY, "erreur": FILL_ORANGE}
    for row_i, (k, v) in enumerate(counts.items(), 2):
        ws2.cell(row=row_i, column=1, value=k).fill = fills[k]
        ws2.cell(row=row_i, column=2, value=v)
        ws2.cell(row=row_i, column=3, value=f"{v/total*100:.0f}%" if total else "0%")

    wb.save(out_path)

# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    if len(sys.argv) < 2:
        print("Usage : python analyse_coherence.py export_produits.xlsx [clé_api]")
        print("        La clé API peut aussi être dans la variable ANTHROPIC_API_KEY")
        sys.exit(1)

    input_path = sys.argv[1]
    api_key    = sys.argv[2] if len(sys.argv) > 2 else os.environ.get("ANTHROPIC_API_KEY", "")

    if not api_key:
        sys.exit("❌ Clé API manquante. Passez-la en argument ou via ANTHROPIC_API_KEY=sk-ant-...")
    if not Path(input_path).exists():
        sys.exit(f"❌ Fichier introuvable : {input_path}")

    out_path = Path(input_path).stem + "_coherence.xlsx"

    print(f"📂 Lecture de {input_path}...")
    products = load_products(input_path)
    print(f"✅ {len(products)} produit(s) avec EAN trouvé(s)\n")

    client  = anthropic.Anthropic(api_key=api_key)
    results = []

    for i, p in enumerate(products, 1):
        print(f"[{i}/{len(products)}] EAN {p['ean']} — {p['title'][:50]}")

        # Téléchargement images
        images_b64 = []
        for url in p["img_urls"][:MAX_IMAGES]:
            b64, ct = fetch_image_b64(url)
            if b64:
                images_b64.append((b64, ct))
                print(f"  ✓ image téléchargée")
            else:
                print(f"  ✗ image inaccessible : {url[:60]}...")

        print(f"  → {len(images_b64)} image(s) envoyée(s) à Claude...")

        try:
            result = analyze(client, p, images_b64)
            vg = result.get("verdict_global", "?")
            emoji = "✅" if vg == "cohérent" else "❌" if vg == "contradiction" else "⚪"
            print(f"  {emoji} Verdict : {vg}")
            results.append({**p, "result": result, "n_images": len(images_b64), "error": ""})
        except Exception as e:
            print(f"  ⚠ Erreur API : {e}")
            results.append({**p, "result": {}, "n_images": len(images_b64), "error": str(e)})

        if i < len(products):
            time.sleep(DELAY_BETWEEN)

    print(f"\n💾 Écriture des résultats dans {out_path}...")
    write_results(results, out_path)

    # Résumé final
    counts = {}
    for r in results:
        vg = r.get("result", {}).get("verdict_global", "erreur" if r.get("error") else "?")
        counts[vg] = counts.get(vg, 0) + 1
    print("\n📊 Résumé :")
    for k, v in sorted(counts.items()):
        print(f"   {k:20s} : {v}")
    print(f"\n✅ Terminé → {out_path}")

if __name__ == "__main__":
    main()
